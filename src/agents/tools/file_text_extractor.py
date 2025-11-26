"""파일에서 텍스트를 추출하는 도구"""
import os
from pathlib import Path
from typing import Optional

# PDF 파싱
import PyPDF2
import pdfplumber

# DOCX 파싱
from docx import Document


def extract_text_from_pdf(file_path: str, use_pdfplumber: bool = True) -> str:
    """
    PDF 파일에서 텍스트 추출
    
    Args:
        file_path: PDF 파일 경로
        use_pdfplumber: True면 pdfplumber 사용, False면 PyPDF2 사용
    
    Returns:
        str: 추출된 텍스트
    """
    text = ""
    
    try:
        if use_pdfplumber:
            # pdfplumber 사용 (더 정확한 텍스트 추출)
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n\n"
        else:
            # PyPDF2 사용 (fallback)
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n\n"
        
        return text.strip()
    
    except Exception as e:
        # pdfplumber 실패 시 PyPDF2로 재시도
        if use_pdfplumber:
            print(f"pdfplumber 추출 실패, PyPDF2로 재시도: {e}")
            return extract_text_from_pdf(file_path, use_pdfplumber = False)
        else:
            raise Exception(f"PDF 텍스트 추출 실패: {e}")


def extract_text_from_docx(file_path: str) -> str:
    """
    DOCX 파일에서 텍스트 추출
    
    Args:
        file_path: DOCX 파일 경로
    
    Returns:
        str: 추출된 텍스트
    """
    try:
        doc = Document(file_path)
        
        # 모든 문단 추출
        paragraphs = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        
        # 표(table) 내용도 추출
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    tables_text.append(row_text)
        
        # 병합
        full_text = "\n\n".join(paragraphs)
        if tables_text:
            full_text += "\n\n[표 내용]\n" + "\n".join(tables_text)
        
        return full_text.strip()
    
    except Exception as e:
        raise Exception(f"DOCX 텍스트 추출 실패: {e}")


def extract_text_from_txt(file_path: str, encoding: str = 'utf-8') -> str:
    """
    TXT 파일에서 텍스트 추출
    
    Args:
        file_path: TXT 파일 경로
        encoding: 파일 인코딩 (기본값: utf-8)
    
    Returns:
        str: 추출된 텍스트
    """
    try:
        with open(file_path, 'r', encoding = encoding) as f:
            return f.read().strip()
    except UnicodeDecodeError:
        # UTF-8 실패 시 cp949 시도 (한국어 환경)
        try:
            with open(file_path, 'r', encoding = 'cp949') as f:
                return f.read().strip()
        except Exception as e:
            raise Exception(f"TXT 파일 인코딩 오류: {e}")
    except Exception as e:
        raise Exception(f"TXT 텍스트 추출 실패: {e}")


def extract_text_from_file(file_path: str) -> str:
    """
    파일 확장자에 따라 적절한 텍스트 추출 메서드 자동 선택
    
    Args:
        file_path: 파일 경로
    
    Returns:
        str: 추출된 텍스트
    
    Raises:
        Exception: 파일이 존재하지 않거나 지원하지 않는 형식일 때
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"파일을 찾을 수 없습니다: {file_path}")
    
    file_ext = Path(file_path).suffix.lower()
    
    if file_ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif file_ext in ['.docx', '.doc']:
        return extract_text_from_docx(file_path)
    elif file_ext == '.txt':
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {file_ext}")


def get_text_preview(text: str, max_length: int = 500) -> str:
    """
    텍스트 미리보기 생성 (디버깅용)
    
    Args:
        text: 전체 텍스트
        max_length: 최대 길이
    
    Returns:
        str: 미리보기 텍스트
    """
    if len(text) <= max_length:
        return text
    return text[:max_length] + "..."
